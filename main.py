import os
import re
import json
import random
import requests
import math
import time
import openpyxl
import csv
from datetime import datetime, timedelta, timezone
import pytz
import ee
from google import genai
from bs4 import BeautifulSoup
from playwright.sync_api import sync_playwright

import urllib3
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY")
MAKE_WEBHOOK_URL = os.environ.get("MAKE_WEBHOOK_URL")
TMD_API_KEY = os.environ.get("TMD_API_KEY")

# ─────────────────────────────────────────────
# Gemini: เว้นระยะระหว่าง call + retry แบบ exponential backoff + jitter
# (ครอบ client เดิม ฟังก์ชันใน ai_brain ที่เรียก client.models.generate_content ได้ผลทันที)
# ─────────────────────────────────────────────
GEMINI_MIN_GAP = float(os.environ.get("GEMINI_MIN_GAP_SEC", "8"))    # เว้นอย่างน้อยกี่วินาทีระหว่าง call
GEMINI_MAX_TRY = int(os.environ.get("GEMINI_MAX_TRY", "5"))

class _ModelsProxy:
    def __init__(self, models):
        self._m = models
        self._last = 0.0

    def __getattr__(self, name):
        return getattr(self._m, name)

    def generate_content(self, *args, **kwargs):
        for i in range(GEMINI_MAX_TRY):
            gap = time.time() - self._last
            if gap < GEMINI_MIN_GAP:
                time.sleep(GEMINI_MIN_GAP - gap)
            try:
                self._last = time.time()
                res = self._m.generate_content(*args, **kwargs)
                self._last = time.time()
                return res
            except Exception as e:
                msg = str(e)
                code = getattr(e, "code", None)
                transient = (code in (429, 500, 503, 504)
                             or "RESOURCE_EXHAUSTED" in msg or "UNAVAILABLE" in msg)
                if "PerDay" in msg or "per day" in msg.lower():   # โควตารายวันหมด รอไปก็ไม่หาย
                    print("❌ Gemini โควตารายวันหมด ไม่ retry")
                    raise
                if not transient or i == GEMINI_MAX_TRY - 1:
                    raise
                wait = min(2 ** (i + 2) + random.uniform(0, 1.5), 60)       # 4,8,16,32,60 + jitter
                m = re.search(r"retry\w*\D{0,12}(\d+(?:\.\d+)?)\s*s", msg, re.I)  # เซิร์ฟเวอร์บอกให้รอกี่วิ
                if m:
                    wait = max(wait, min(float(m.group(1)) + 1, 90))
                print(f"⚠️ Gemini {code or 'transient'} → รอ {wait:.0f}s แล้วลองใหม่ ({i+1}/{GEMINI_MAX_TRY})")
                time.sleep(wait)

class _ClientProxy:
    def __init__(self, c):
        self._c = c
        self.models = _ModelsProxy(c.models)

    def __getattr__(self, name):
        return getattr(self._c, name)

client = _ClientProxy(genai.Client(api_key=GEMINI_API_KEY))

tz = pytz.timezone('Asia/Bangkok')
now = datetime.now(tz)
THAI_MONTHS = ["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
               "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]
THAI_DAYS   = ["จันทร์", "อังคาร", "พุธ", "พฤหัสบดี", "ศุกร์", "เสาร์", "อาทิตย์"]
thai_month       = THAI_MONTHS[now.month - 1]
thai_year        = now.year + 543
thai_day_of_week = THAI_DAYS[now.weekday()]   # 0=จันทร์ … 6=อาทิตย์
date_str         = f"{now.day} {thai_month} {thai_year}"
time_str         = now.strftime("%H:%M น.")

INBURI_LAT = 14.9961
INBURI_LON = 100.3253

# ─────────────────────────────────────────────
# State
# ─────────────────────────────────────────────
STATE_FILE = "state.json"

# cache ข้ามรอบ เก็บใน state.json  {key: {"data":..., "ts": epoch}}
_CACHE = {}
CACHE_MAX_AGE = {                      # อายุสูงสุดที่ยอมใช้ข้อมูลเก่า (วินาที)
    "tmd_obs": 3 * 3600, "tmd_nwp": 2 * 3600, "tomorrow": 2 * 3600,
    "weather_om": 3 * 3600, "wl": 3 * 3600, "dam_discharge": 6 * 3600,
}
_STALE = {}                            # key -> "HH:MM" ของข้อมูลที่ต้องใช้ cache เก่าในรอบนี้

def cache_put(key, data):
    _CACHE[key] = {"data": data, "ts": time.time()}

def cache_get(key, max_age=None):
    """คืน (data, 'HH:MM') ถ้ามี cache และอายุไม่เกินกำหนด ไม่งั้น (None, None)"""
    e = _CACHE.get(key)
    if not isinstance(e, dict) or "data" not in e:
        return None, None
    base = key.split(":")[0]
    limit = max_age if max_age is not None else CACHE_MAX_AGE.get(base, 3 * 3600)
    try:
        ts = float(e.get("ts", 0))
    except Exception:
        return None, None
    if time.time() - ts > limit:
        return None, None
    return e["data"], datetime.fromtimestamp(ts, tz).strftime("%H:%M")

def mark_stale(key, hhmm):
    _STALE[key] = hhmm

def load_state() -> dict:
    state = {}
    try:
        with open(STATE_FILE, 'r', encoding='utf-8') as f:
            state = json.load(f)
        if not isinstance(state, dict):
            state = {}
    except FileNotFoundError:
        pass
    except Exception as e:
        print(f"⚠️ อ่าน state.json ไม่ได้ ({e}) เริ่มจากค่าว่าง")
    _CACHE.clear()
    c = state.get("cache")
    if isinstance(c, dict):
        _CACHE.update(c)
    return state

def save_state(state: dict):
    """เขียนแบบ atomic: เขียนไฟล์ชั่วคราว → fsync → os.replace (job ถูก kill กลางคันก็ไม่ทำไฟล์เดิมพัง)"""
    tmp = STATE_FILE + ".tmp"
    try:
        out = dict(state)
        out["cache"] = dict(_CACHE)
        with open(tmp, 'w', encoding='utf-8') as f:
            json.dump(out, f, ensure_ascii=False, indent=2)
            f.flush()
            os.fsync(f.fileno())
        os.replace(tmp, STATE_FILE)
    except Exception as e:
        print(f"⚠️ บันทึก state ไม่ได้: {e}")
        try:
            os.remove(tmp)
        except OSError:
            pass

def build_compare_text(current, previous, unit: str, label: str) -> str:
    if previous is None or current is None:
        return ""
    try:
        delta = round(float(current) - float(previous), 2)
        prev_str = f"{float(previous):.2f}".rstrip('0').rstrip('.')
        if delta > 0:
            return f"⬆️ เพิ่มขึ้น {abs(delta):.2f} {unit} จากเมื่อวาน ({prev_str} {unit})"
        elif delta < 0:
            return f"⬇️ ลดลง {abs(delta):.2f} {unit} จากเมื่อวาน ({prev_str} {unit})"
        else:
            return f"➡️ ทรงตัว เท่ากับเมื่อวาน ({prev_str} {unit})"
    except Exception:
        return ""

# ─────────────────────────────────────────────
# ฟังก์ชันระยะทาง (เดิม)
# ─────────────────────────────────────────────
def get_dist(lat1, lon1, lat2, lon2):
    R = 6371
    dlat = math.radians(float(lat2) - float(lat1))
    dlon = math.radians(float(lon2) - float(lon1))
    a = (math.sin(dlat/2)**2
         + math.cos(math.radians(float(lat1)))
         * math.cos(math.radians(float(lat2)))
         * math.sin(dlon/2)**2)
    return R * 2 * math.asin(math.sqrt(a))

# ─────────────────────────────────────────────
# HTTP helper: retry + backoff (429/5xx/timeout) และแคช Tomorrow.io ต่อรอบรัน
# ─────────────────────────────────────────────
def _backoff_wait(i, retry_after=None):
    w = min(2 ** i + random.uniform(0, 1.5), 30)        # 1-2.5s, 2-3.5s, 4-5.5s ...
    if retry_after:
        w = max(w, min(retry_after, 60))
    return w

def _http_get_json(url, timeout=20, retries=3, headers=None):
    for i in range(retries):
        retry_after = None
        try:
            r = requests.get(url, timeout=timeout, headers=headers)
            if r.status_code == 200:
                return r.json()
            if r.status_code not in (429, 500, 502, 503, 504):
                print(f"⚠️ HTTP {r.status_code}: {url.split('?')[0]}")
                return None
            try:
                retry_after = float(r.headers.get("Retry-After", ""))
            except ValueError:
                pass
            print(f"⚠️ HTTP {r.status_code} ({i+1}/{retries}): {url.split('?')[0]}")
        except Exception as e:
            print(f"⚠️ GET ล้มเหลว ({i+1}/{retries}) {url.split('?')[0]}: {e}")
        if i < retries - 1:
            time.sleep(_backoff_wait(i, retry_after))
    return None

def _http_get_raw(url, headers=None, timeout=15, retries=3):
    """GET พร้อม retry (429/5xx/timeout) คืน Response ตัวสุดท้าย หรือ None ถ้าเชื่อมต่อไม่ได้เลย
    401/403/404 ฯลฯ ไม่ retry (ลองซ้ำก็ไม่หาย) ให้ผู้เรียกจัดการเอง"""
    res = None
    for i in range(retries):
        retry_after = None
        try:
            res = requests.get(url, headers=headers, timeout=timeout)
            if res.status_code not in (429, 500, 502, 503, 504):
                return res
            try:
                retry_after = float(res.headers.get("Retry-After", ""))
            except ValueError:
                pass
            print(f"⚠️ HTTP {res.status_code} ({i+1}/{retries}): {url.split('?')[0]}")
        except Exception as e:
            print(f"⚠️ GET ล้มเหลว ({i+1}/{retries}) {url.split('?')[0]}: {e}")
        if i < retries - 1:
            time.sleep(_backoff_wait(i, retry_after))
    return res

def _as_list(x):
    """API บางตัวคืน dict เดี่ยวเมื่อมีรายการเดียว / null เมื่อไม่มี → ทำให้เป็น list เสมอ"""
    if x is None:
        return []
    return x if isinstance(x, list) else [x]

def _http_get_text(url, timeout=20, retries=3):
    for i in range(retries):
        try:
            r = requests.get(url, timeout=timeout)
            if r.status_code == 200:
                return r.text
            print(f"⚠️ HTTP {r.status_code} ({i+1}/{retries}): {url.split('?')[0]}")
        except Exception as e:
            print(f"⚠️ GET ล้มเหลว ({i+1}/{retries}) {url.split('?')[0]}: {e}")
        if i < retries - 1:
            time.sleep(_backoff_wait(i))
    return None

_TMR_CACHE = {}
def _tomorrow_forecast():
    """เรียก Tomorrow.io ครั้งเดียวต่อรอบรัน ใช้ร่วมกันทั้ง get_weather และ get_rain_storm_forecast
    เรียกไม่สำเร็จ → ใช้ cache ใน state.json (อายุ ≤ 2 ชม.) แล้วตัดชั่วโมงที่ผ่านไปแล้วทิ้ง"""
    key = os.environ.get("TOMORROW_API_KEY")
    if not key:
        return None
    if "data" in _TMR_CACHE:
        return _TMR_CACHE["data"]

    data = _http_get_json(
        f"https://api.tomorrow.io/v4/weather/forecast"
        f"?location={INBURI_LAT},{INBURI_LON}&apikey={key}", timeout=15, retries=3)
    slim = None
    if data:
        try:
            tl = data.get('timelines', {})
            slim = {'timelines': {'minutely': (tl.get('minutely') or [])[:30],
                                  'hourly':   (tl.get('hourly') or [])[:24]}}
            if not slim['timelines']['hourly']:
                slim = None
        except Exception:
            slim = None
    if slim:
        cache_put("tomorrow", slim)
        _TMR_CACHE["data"] = slim
        return slim

    cached, ts = cache_get("tomorrow")
    if cached:
        cutoff = datetime.now(timezone.utc) - timedelta(hours=1)
        hourly = []
        for h in cached['timelines']['hourly']:
            try:
                if datetime.fromisoformat(h['time'].replace('Z', '+00:00')) >= cutoff:
                    hourly.append(h)
            except Exception:
                continue
        if hourly:
            mark_stale("tomorrow", ts)
            print(f"♻️ Tomorrow.io ดึงสดไม่ได้ → ใช้ cache ณ {ts} น.")
            _TMR_CACHE["data"] = {'timelines': {'minutely': [], 'hourly': hourly}}   # minutely หมดอายุเร็ว ทิ้ง
            return _TMR_CACHE["data"]
    _TMR_CACHE["data"] = None
    return None

# ─────────────────────────────────────────────
# TMD ชุดที่ 1: ผลตรวจวัดและพยากรณ์อากาศ (Observation)
# ─────────────────────────────────────────────
def _get_tmd_observation_live() -> dict:
    result = {'available': False}
    if not TMD_API_KEY:
        print("⚠️ TMD_API_KEY ไม่พบ")
        return result
    try:
        url = (
            "https://data.tmd.go.th/api/Weather3Hours/v2/"
            f"?APIkey={TMD_API_KEY}&station_type=ตรวจอากาศผิวพื้น"
        )
        headers = {'Accept': 'application/json'}
        res = _http_get_raw(url, headers=headers, timeout=15, retries=3)
        if res is None:
            print("⚠️ TMD Obs: เชื่อมต่อไม่ได้")
            return result
        print(f"TMD Observation HTTP: {res.status_code}")
        if res.status_code in (401, 403):
            print("⚠️ TMD Obs: token ไม่ถูกต้อง/หมดอายุ → ตรวจ/ต่ออายุ TMD_API_KEY")
            return result
        if res.status_code != 200:
            return result

        raw = res.text.strip()
        if not raw:
            print("⚠️ TMD Obs: response body ว่างเปล่า (ยังไม่มีข้อมูลในรอบนี้)")
            return result
        if raw.startswith('<'):
            print("⚠️ TMD Obs: ได้ HTML กลับมาแทน JSON")
            return result
        try:
            data = res.json()
        except Exception as je:
            print(f"⚠️ TMD Obs: JSON parse ไม่ได้ ({je})")
            return result

        stations = [st for st in _as_list((data.get('Stations') or {}).get('Station')) if isinstance(st, dict)]

        def _parse_cands(max_dist_km):
            cands = []
            for st in stations:
                try:
                    lat  = float(st.get('Latitude', 0))
                    lon  = float(st.get('Longitude', 0))
                    dist = get_dist(INBURI_LAT, INBURI_LON, lat, lon)
                    if dist > max_dist_km:
                        continue
                    obs      = st.get('Observation', {})
                    rain_raw = obs.get('Rainfall', {})
                    rain     = float(rain_raw.get('Value') or 0) if isinstance(rain_raw, dict) else float(rain_raw or 0)
                    temp     = float((obs.get('AirTemperature', {}) or {}).get('Value') or 0)
                    hum      = float((obs.get('RelativeHumidity', {}) or {}).get('Value') or 0)
                    wind     = float((obs.get('WindSpeed', {}) or {}).get('Value') or 0)
                    cands.append({
                        'name': st.get('StationNameThai', st.get('StationNameEng', '?')),
                        'dist': dist, 'rain_3h': rain,
                        'temp': temp, 'humidity': hum, 'wind_speed': wind,
                    })
                except Exception:
                    continue
            return cands

        candidates = _parse_cands(35)
        if not candidates:
            candidates = _parse_cands(80)
            print("⚠️ TMD Obs: ไม่พบสถานีใน 35 กม. ขยายเป็น 80 กม.")
        if not candidates:
            print("⚠️ TMD Obs: ไม่พบสถานีเลย")
            return result

        candidates.sort(key=lambda x: x['dist'])
        top3    = candidates[:3]
        total_w = sum(1.0 / max(c['dist'], 0.1) ** 2 for c in top3)
        w_rain  = sum(c['rain_3h']    / max(c['dist'], 0.1) ** 2 for c in top3) / total_w
        w_temp  = sum(c['temp']       / max(c['dist'], 0.1) ** 2 for c in top3) / total_w
        w_hum   = sum(c['humidity']   / max(c['dist'], 0.1) ** 2 for c in top3) / total_w
        w_wind  = sum(c['wind_speed'] / max(c['dist'], 0.1) ** 2 for c in top3) / total_w
        best    = candidates[0]
        rain_report = max(w_rain, best['rain_3h'])

        result.update({
            'available': True,
            'rain_3h': round(rain_report, 2),
            'temp': round(w_temp, 1),
            'humidity': round(w_hum, 1),
            'wind_speed': round(w_wind, 1),
            'station_name': best['name'],
            'dist_km': round(best['dist'], 1),
        })
        print(f"✅ TMD Obs ({len(top3)} สถานี): {best['name']} {best['dist']:.1f}กม. ฝน={rain_report:.2f}มม.")
    except Exception as e:
        print(f"⚠️ TMD Observation error: {e}")
    return result


def get_tmd_observation() -> dict:
    """ดึงสด → สำเร็จก็เก็บ cache | ล้มเหลว → ใช้ cache ≤ 3 ชม. พร้อม stale_ts"""
    r = _get_tmd_observation_live()
    if r.get('available'):
        cache_put("tmd_obs", r)
        return r
    cached, ts = cache_get("tmd_obs")
    if isinstance(cached, dict) and cached.get('available'):
        mark_stale("tmd_obs", ts)
        print(f"♻️ TMD Obs ดึงสดไม่ได้ → ใช้ cache ณ {ts} น.")
        return {**cached, 'stale_ts': ts}
    return r


# ─────────────────────────────────────────────
# TMD ชุดที่ 2: พยากรณ์จากกรมอุตุฯ (WeatherForecast)
# ─────────────────────────────────────────────
def _get_tmd_nwp_forecast_live() -> dict:
    result = {'available': False}
    if not TMD_API_KEY:
        return result

    endpoints = [
        (
            "https://data.tmd.go.th/nwpapi/v1/forecast/location/hourly/at"
            f"?lat={INBURI_LAT}&lon={INBURI_LON}&fields=tc,rh,rain,ws10m&duration=6"
        ),
        (
            "https://data.tmd.go.th/api/WeatherForecast/v2/"
            f"?APIkey={TMD_API_KEY}"
            f"&lat={INBURI_LAT}&lon={INBURI_LON}"
        ),
        (
            "https://data.tmd.go.th/api/NowcastForecast/v2/"
            f"?APIkey={TMD_API_KEY}"
            f"&lat={INBURI_LAT}&lon={INBURI_LON}"
            f"&fields=rain,tc,rh,ws,thunderstorm"
        ),
    ]

    data = None
    for ep_url in endpoints:
        try:
            headers = {'Accept': 'application/json'}
            if "/nwpapi/" in ep_url:
                headers['Authorization'] = f"Bearer {TMD_API_KEY}"
            res = _http_get_raw(ep_url, headers=headers, timeout=15, retries=2)
            if res is None:
                print("⚠️ TMD NWP: เชื่อมต่อ endpoint ไม่ได้")
                continue
            ep_name = ("nwpapi-v1" if "/nwpapi/" in ep_url else
                       "WeatherForecast" if "WeatherForecast" in ep_url else "NowcastForecast")
            print(f"TMD NWP ({ep_name}) HTTP: {res.status_code}")
            if res.status_code in (401, 403):
                print(f"⚠️ TMD NWP ({ep_name}): token ไม่ถูกต้อง/หมดอายุ")
                continue
            raw = res.text.strip()
            if not raw:
                print(f"⚠️ TMD NWP ({ep_name}): body ว่างเปล่า")
                continue
            if raw.startswith('<'):
                print(f"⚠️ TMD NWP ({ep_name}): ได้ HTML กลับมา (endpoint ไม่รองรับ)")
                continue
            if res.status_code == 200:
                data = res.json()
                print(f"✅ TMD NWP: ใช้ {ep_name}")
                break
        except Exception as ep_e:
            print(f"⚠️ TMD NWP endpoint error: {ep_e}")
            continue

    if data is None:
        print("⚠️ TMD NWP: ทุก endpoint ล้มเหลว")
        return result

    try:
        wf_list   = _as_list(data.get('WeatherForecasts'))
        first     = wf_list[0] if wf_list and isinstance(wf_list[0], dict) else {}
        forecasts = (
            _as_list(first.get('forecasts'))
            or _as_list(data.get('forecasts'))
            or _as_list(data.get('Forecasts'))
            or wf_list
        )
        forecasts = [f for f in forecasts if isinstance(f, dict)]

        next_6h = forecasts[:6]
        if not next_6h:
            print("⚠️ TMD NWP: ไม่พบข้อมูลพยากรณ์ (keys:", list(data.keys()), ")")
            return result

        rain_vals, thunder_vals, wind_vals, temp_vals = [], [], [], []
        for fc in next_6h:
            d = fc.get('data', fc)
            try: rain_vals.append(float(d.get('rain', 0) or 0))
            except Exception: pass
            try: thunder_vals.append(float(d.get('thunderstorm', 0) or 0))
            except Exception: pass
            try: wind_vals.append(float(d.get('ws', d.get('ws10m', 0)) or 0))
            except Exception: pass
            try: temp_vals.append(float(d.get('tc', 0) or 0))
            except Exception: pass

        max_rain    = max(rain_vals,    default=0)
        max_thunder = max(thunder_vals, default=0)
        max_wind    = max(wind_vals,    default=0)
        avg_temp    = sum(temp_vals) / len(temp_vals) if temp_vals else 0

        if max_rain >= 35 or (max_rain >= 20 and max_thunder >= 50):
            desc = f"🚨 NWP กรมอุตุฯ: ฝนหนักมาก {max_rain:.1f} มม./ชม. ฟ้าคะนอง {max_thunder:.0f}%"
        elif max_rain >= 10 or (max_rain >= 5 and max_thunder >= 30):
            desc = f"⚠️ NWP กรมอุตุฯ: ฝนหนัก {max_rain:.1f} มม./ชม. ฟ้าคะนอง {max_thunder:.0f}%"
        elif max_rain >= 1:
            desc = f"🌧️ NWP กรมอุตุฯ: มีฝน {max_rain:.1f} มม./ชม. ใน 6 ชม. ข้างหน้า"
        else:
            desc = f"☀️ NWP กรมอุตุฯ: ไม่มีฝนใน 6 ชม. ข้างหน้า"

        result.update({
            'available': True,
            'rain_1h_max': max_rain,
            'thunder_max': max_thunder,
            'wind_max': max_wind,
            'temp_avg': round(avg_temp, 1),
            'description': desc,
        })
        print(f"✅ TMD NWP: ฝน={max_rain:.1f} ฟ้าคะนอง={max_thunder:.0f}% ลม={max_wind:.1f}")
    except Exception as e:
        print(f"⚠️ TMD NWP parse error: {e}")
    return result


def get_tmd_nwp_forecast() -> dict:
    r = _get_tmd_nwp_forecast_live()
    if r.get('available'):
        cache_put("tmd_nwp", r)
        return r
    cached, ts = cache_get("tmd_nwp")
    if isinstance(cached, dict) and cached.get('available'):
        mark_stale("tmd_nwp", ts)
        print(f"♻️ TMD NWP ดึงสดไม่ได้ → ใช้ cache ณ {ts} น.")
        return {**cached, 'stale_ts': ts}
    return r


# ─────────────────────────────────────────────
# Open-Meteo: ฝนที่ตกจริงย้อนหลัง (backup)
# ─────────────────────────────────────────────
def get_actual_rain_last_hour() -> dict:
    result = {'rain_1h': None, 'rain_3h': None}
    try:
        url = (
            f"https://api.open-meteo.com/v1/forecast"
            f"?latitude={INBURI_LAT}&longitude={INBURI_LON}"
            f"&hourly=precipitation&past_hours=3&forecast_hours=1"
            f"&timezone=Asia%2FBangkok"
        )
        res    = _http_get_json(url, timeout=25, retries=3)
        if not res:
            return result
        times  = res['hourly']['time']
        precip = res['hourly']['precipitation']

        now_str = datetime.now(tz).strftime("%Y-%m-%dT%H:00")
        current_idx = next((i for i, t in enumerate(times) if t == now_str), None)
        if current_idx is None:
            past = [i for i, t in enumerate(times) if t <= now_str]
            current_idx = max(past) if past else 0

        if current_idx >= 1:
            result['rain_1h'] = round(float(precip[current_idx - 1] or 0), 2)
        if current_idx >= 3:
            result['rain_3h'] = round(
                sum(float(p or 0) for p in precip[max(0, current_idx-3):current_idx]), 2
            )
    except Exception as e:
        print(f"⚠️ Open-Meteo actual rain error: {e}")
    return result

# ─────────────────────────────────────────────
# Tomorrow.io: minutely + 6h forecast
# ─────────────────────────────────────────────
def get_rain_storm_forecast() -> dict:
    TOMORROW_API_KEY = os.environ.get("TOMORROW_API_KEY")
    if not TOMORROW_API_KEY:
        return {}
    try:
        data = _tomorrow_forecast()
        if not data:
            return {}

        hourly_6h   = data['timelines']['hourly'][:6]
        hourly_24h  = data['timelines']['hourly'][:24]
        minutely_30 = data['timelines'].get('minutely', [])[:30]

        rain_now_prob, rain_now_intensity = 0, 0
        if minutely_30:
            rain_now_prob = max(
                float(m['values'].get('precipitationProbability', 0) or 0) for m in minutely_30)
            rain_now_intensity = max(
                float(m['values'].get('precipitationIntensity', 0) or 0) for m in minutely_30)

        danger_slots, warn_slots, watch_slots, heavy_slots = [], [], [], []
        for h in hourly_6h:
            v         = h.get('values', {})
            rain_prob = float(v.get('precipitationProbability', 0) or 0)
            thunder   = float(v.get('thunderstormProbability',  0) or 0)
            wind_spd  = float(v.get('windSpeed', 0) or 0)
            wind_gust = float(v.get('windGust',  0) or 0)
            intensity = float(v.get('precipitationIntensity', 0) or 0)
            try:
                hour_label = datetime.fromisoformat(
                    h['time'].replace('Z', '+00:00')).astimezone(tz).strftime("%H:%M")
            except Exception:
                hour_label = "?"

            slot     = {'time': hour_label, 'rain_prob': rain_prob, 'thunder': thunder,
                        'wind': wind_spd, 'gust': wind_gust, 'intensity': intensity}
            eff_wind = max(wind_spd, wind_gust * 0.7)

            if   rain_prob >= 80 and (thunder >= 50 or eff_wind >= 14): danger_slots.append(slot)
            elif rain_prob >= 60 and (thunder >= 30 or eff_wind >= 10): warn_slots.append(slot)
            elif rain_prob >= 40 and (thunder >= 20 or eff_wind >= 7):  watch_slots.append(slot)
            elif rain_prob >= 60:                                         heavy_slots.append(slot)

        now_summary = ""
        if rain_now_prob >= 70 or rain_now_intensity >= 2.0:
            now_summary = f"🌧️ ฝนกำลังตกอยู่ขณะนี้ ({rain_now_intensity:.1f} มม./ชม.)"
        elif rain_now_prob >= 40:
            now_summary = f"🌦️ มีโอกาสฝนตกใน 30 นาทีนี้ ({rain_now_prob:.0f}%)"

        max_rain_24h = max(
            (float(h['values'].get('precipitationProbability', 0) or 0) for h in hourly_24h),
            default=0)

        return {
            'danger_slots': danger_slots, 'warn_slots': warn_slots,
            'watch_slots': watch_slots,   'heavy_slots': heavy_slots,
            'now_summary': now_summary,
            'rain_now_prob': rain_now_prob, 'rain_now_intensity': rain_now_intensity,
            'max_rain_24h': max_rain_24h,
        }
    except Exception as e:
        print(f"⚠️ Tomorrow.io error: {e}")
        return {}


# ─────────────────────────────────────────────
# รวมข้อมูลฝนทั้ง 4 แหล่ง → สรุปเดียว
# ─────────────────────────────────────────────
def get_comprehensive_rain_info() -> dict:
    print("กำลังดึงข้อมูลฝน/พายุจาก 4 แหล่ง...")

    tmd_obs = get_tmd_observation()        # แหล่ง 1: ตรวจวัดจริง
    tmd_nwp = get_tmd_nwp_forecast()       # แหล่ง 2: NWP กรมอุตุฯ
    tmr     = get_rain_storm_forecast()    # แหล่ง 3: Tomorrow.io minutely+6h
    actual  = get_actual_rain_last_hour()  # แหล่ง 4: Open-Meteo ย้อนหลัง

    sources_used = []
    if tmd_obs.get('available'):            sources_used.append('TMD สถานี')
    if tmd_nwp.get('available'):            sources_used.append('TMD NWP')
    if tmr:                                 sources_used.append('Tomorrow.io')
    if actual.get('rain_1h') is not None:   sources_used.append('Open-Meteo')

    # ── ฝนจริงที่ตกแล้ว ──────────────────────
    rain_1h     = actual.get('rain_1h', 0) or 0
    rain_3h     = actual.get('rain_3h', 0) or 0
    tmd_rain_3h = tmd_obs.get('rain_3h', 0) or 0

    actual_rain_best = tmd_rain_3h if (tmd_obs.get('available') and tmd_rain_3h > 0) else rain_3h

    actual_text = ""
    if tmd_obs.get('available') and tmd_rain_3h >= 1:
        actual_text = (
            f"🌧️ สถานีอุตุฯ {tmd_obs['station_name']} "
            f"(ห่าง {tmd_obs['dist_km']} กม.): ฝน 3 ชม. = {tmd_rain_3h:.1f} มม."
            + (f" (ข้อมูล ณ {tmd_obs['stale_ts']} น.)" if tmd_obs.get('stale_ts') else "")
        )
    elif rain_1h >= 2:
        actual_text = f"🌧️ มีฝนตกในชั่วโมงที่ผ่านมา {rain_1h:.1f} มม."

    # ── NWP กรมอุตุฯ ─────────────────────────
    nwp_text = tmd_nwp.get('description', '') if tmd_nwp.get('available') else ''
    if nwp_text and tmd_nwp.get('stale_ts'):
        nwp_text += f" (ข้อมูล ณ {tmd_nwp['stale_ts']} น.)"

    # ── Tomorrow.io ──────────────────────────
    now_text      = tmr.get('now_summary', '') if tmr else ''
    tmr_slot_text = ''
    if tmr:
        if tmr.get('danger_slots'):
            s = tmr['danger_slots'][0]
            times = ", ".join(x['time'] for x in tmr['danger_slots'][:3])
            tmr_slot_text = (f"🚨 Tomorrow.io: เสี่ยงพายุรุนแรงช่วง {times} น. "
                             f"โอกาสฝน {s['rain_prob']:.0f}% ลม {s['wind']:.1f} m/s")
        elif tmr.get('warn_slots'):
            s = tmr['warn_slots'][0]
            times = ", ".join(x['time'] for x in tmr['warn_slots'][:3])
            tmr_slot_text = (f"⚠️ Tomorrow.io: เสี่ยงพายุช่วง {times} น. "
                             f"โอกาสฝน {s['rain_prob']:.0f}%")
        elif tmr.get('watch_slots'):
            s = tmr['watch_slots'][0]
            times = ", ".join(x['time'] for x in tmr['watch_slots'][:3])
            tmr_slot_text = f"👀 Tomorrow.io: เฝ้าระวังช่วง {times} น. โอกาสฝน {s['rain_prob']:.0f}%"

    if tmr_slot_text and _STALE.get("tomorrow"):
        tmr_slot_text += f" (ข้อมูล ณ {_STALE['tomorrow']} น.)"

    # ── รวม ──────────────────────────────────
    parts = [p for p in [actual_text, nwp_text, now_text, tmr_slot_text] if p]
    if not parts:
        max_r = tmr.get('max_rain_24h', 0) if tmr else 0
        parts = ([f"อาจมีฝนประปราย โอกาสสูงสุดวันนี้ {max_r:.0f}%"]
                 if max_r >= 30 else
                 ["ท้องฟ้าโปร่งดี ไม่มีพายุในพื้นที่อินทร์บุรี"])

    combined   = " | ".join(parts)
    confidence = f"(ข้อมูลจาก: {', '.join(sources_used)})" if sources_used else ""

    # ── ระดับความเสี่ยงรวม ────────────────────
    nwp_rain    = tmd_nwp.get('rain_1h_max', 0) or 0
    nwp_thunder = tmd_nwp.get('thunder_max', 0) or 0
    has_danger  = (bool(tmr.get('danger_slots')) or nwp_rain >= 35
                   or (nwp_rain >= 20 and nwp_thunder >= 50) or actual_rain_best >= 15)
    has_warn    = (bool(tmr.get('warn_slots'))   or nwp_rain >= 10 or actual_rain_best >= 5)
    has_watch   = (bool(tmr.get('watch_slots'))  or nwp_rain >= 1  or actual_rain_best >= 1)

    risk_level = ('danger' if has_danger else
                  'warn'   if has_warn   else
                  'watch'  if has_watch  else 'none')

    return {
        'summary': combined, 'confidence': confidence,
        'risk_level': risk_level, 'sources': sources_used,
        'tmd_obs': tmd_obs, 'tmd_nwp': tmd_nwp,
        'forecast': tmr,    'actual_rain_1h': rain_1h,
    }


# ─────────────────────────────────────────────
# ฟังก์ชันเดิม ไม่แตะ
# ─────────────────────────────────────────────
def get_hotspots():
    EE_JSON_KEY = os.environ.get("EE_JSON_KEY")
    if not EE_JSON_KEY:
        return "N/A"
    try:
        json_key    = json.loads(EE_JSON_KEY)
        credentials = ee.ServiceAccountCredentials(json_key['client_email'], key_data=EE_JSON_KEY)
        ee.Initialize(credentials)
        inburi_area = ee.Geometry.Point([100.3273, 15.0076]).buffer(12000)
        end_date    = ee.Date(datetime.now(tz))
        start_date  = end_date.advance(-24, 'hour')
        dataset = (ee.ImageCollection("FIRMS")
                   .filterBounds(inburi_area)
                   .filterDate(start_date, end_date))
        if dataset.size().getInfo() == 0:
            return 0
        fires        = dataset.select('T21').max()
        fires_masked = fires.updateMask(fires.gt(0))
        stats = fires_masked.reduceRegion(
            reducer=ee.Reducer.count(), geometry=inburi_area,
            scale=1000, maxPixels=1e9)
        count = stats.get('T21').getInfo()
        return int(count) if count is not None else 0
    except Exception as e:
        print(f"⚠️ ระบบดาวเทียมขัดข้อง: {e}")
        return "N/A"

def _safe_float(v):
    try:
        return float(str(v).replace(',', '').strip())
    except Exception:
        return None

def _parse_air4thai_age_seconds(st):
    patterns = ["%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S",
                "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S"]
    try:
        lu        = st.get('LastUpdate', {}) if isinstance(st, dict) else {}
        date_text = str(lu.get('date') or st.get('date') or '').strip()
        time_text = str(lu.get('time') or st.get('time') or '').strip()
        if not date_text or not time_text:
            return None
        raw = f"{date_text} {time_text}"
        dt  = None
        for fmt in patterns:
            try:
                dt = datetime.strptime(raw, fmt)
                break
            except Exception:
                pass
        if dt is None:
            return None
        if dt.tzinfo is None:
            dt = tz.localize(dt)
        return max(0, int((datetime.now(tz) - dt).total_seconds()))
    except Exception:
        return None

def _weighted_pm25(rows):
    total_w, total_v = 0.0, 0.0
    for r in rows:
        w = (1 / ((r['distance'] + 1.0) ** 1.6)
             * max(0.15, 1 - (r['age'] / max(r['max_age'], 1)))
             * (1.0 if r['source'] == 'air4thai' else 0.85))
        total_w += w
        total_v += r['pm25'] * w
    return total_v / total_w if total_w else None


def classify_pm25_th(pm25_value):
    pm = _safe_float(pm25_value)
    if pm is None:
        return {
            'label': 'ไม่ทราบ',
            'brief': 'เซ็นเซอร์ขัดข้อง',
            'warning': 'ให้แจ้งตามตรงว่าระบบดึงค่าฝุ่นไม่ได้ และเตือนให้ป้องกันตัวไว้ก่อน',
        }

    if pm <= 15:
        return {
            'label': 'ดีมาก',
            'brief': 'อากาศดีมาก',
            'warning': 'พูดได้ว่าอากาศดีมาก',
        }
    if pm <= 25:
        return {
            'label': 'ดี',
            'brief': 'อากาศค่อนข้างดี',
            'warning': 'พูดได้ว่าอากาศดี แต่ไม่ต้องเว่อร์',
        }
    if pm <= 37.5:
        return {
            'label': 'ปานกลาง',
            'brief': 'เริ่มมีฝุ่นและกลุ่มเสี่ยงควรระวัง',
            'warning': 'ห้ามเขียนว่าอากาศดี ให้สื่อว่าเริ่มมีฝุ่น/กลุ่มเสี่ยงควรระวัง',
        }
    if pm <= 75:
        return {
            'label': 'เริ่มมีผลกระทบต่อสุขภาพ',
            'brief': 'ฝุ่นเยอะ มีผลต่อสุขภาพ',
            'warning': 'ห้ามเขียนว่าอากาศดี ให้เตือนเรื่องหน้ากากและลดกิจกรรมกลางแจ้ง',
        }
    return {
        'label': 'มีผลกระทบต่อสุขภาพมาก',
        'brief': 'ฝุ่นหนักมาก อันตราย',
        'warning': 'ห้ามเขียนว่าอากาศดีเด็ดขาด ให้เตือนจริงจังให้อยู่ในอาคารและใส่หน้ากาก',
    }


def build_pm25_instruction(pm25_value, source_label='ระบบคัดกรองฝุ่น'):
    info = classify_pm25_th(pm25_value)
    pm = _safe_float(pm25_value)
    if pm is None:
        return ("เซ็นเซอร์ขัดข้อง ดึงค่าไม่ได้ ให้เขียนบอกลูกเพจไปตามตรงว่าระบบขัดข้อง "
                "แต่ให้เตือนว่าควรใส่หน้ากากอนามัยป้องกันไว้ก่อนเพื่อความปลอดภัย")

    return (
        f"ค่าฝุ่นอยู่ที่ {pm:.1f} µg/m³ จาก {source_label} "
        f"จัดอยู่ระดับ '{info['label']}' ({info['brief']}) "
        f"{info['warning']} และให้ใช้ถ้อยคำสอดคล้องกับระดับนี้เท่านั้น"
    )


def get_accurate_pm25(return_meta=False):
    PM_LAT, PM_LON = 15.0076, 100.3273
    STRICT_KM = 20
    WIDE_KM = 50
    STRICT_AGE = 3600
    WIDE_AGE = 10800
    headers = {'User-Agent': 'Mozilla/5.0'}

    air4thai_rows = []
    gistda_value = waqi_value = owm_value = openmeteo_value = None
    selected_source = None

    try:
        res = requests.get(
            f"https://pm25.gistda.or.th/rest/getPM25byLocation"
            f"?lat={PM_LAT}&lng={PM_LON}&t={int(time.time())}",
            headers=headers, timeout=15, verify=False)
        print(f"GISTDA HTTP: {res.status_code}")
        if res.status_code == 200:
            payload = res.json()
            pm = _safe_float((payload.get('data', payload)).get('pm25'))
            print(f"GISTDA PM2.5: {pm}")
            if pm is not None:
                gistda_value = pm
    except Exception as e:
        print(f"⚠️ GISTDA error: {e}")

    try:
        res = requests.get(
            f"http://air4thai.pcd.go.th/services/getNewAQI_JSON.php?t={int(time.time())}",
            headers=headers, timeout=15, verify=False)
        print(f"Air4Thai HTTP: {res.status_code}")
        if res.status_code == 200:
            nearby = []
            for st in res.json().get('stations', []):
                pm25_val = _safe_float(st.get('LastUpdate', {}).get('PM25', {}).get('value'))
                if pm25_val is None:
                    continue
                lat = _safe_float(st.get('lat'))
                lon = _safe_float(st.get('long'))
                if lat is None or lon is None:
                    continue
                dist = get_dist(PM_LAT, PM_LON, lat, lon)
                if dist > WIDE_KM:
                    continue
                age = _parse_air4thai_age_seconds(st)
                if age is None:
                    age = WIDE_AGE + 1
                name = st.get('nameTH') or st.get('nameEN') or '?'
                nearby.append((dist, age, name, pm25_val))
                air4thai_rows.append({
                    'source': 'air4thai',
                    'pm25': pm25_val,
                    'distance': dist,
                    'age': age,
                    'max_age': WIDE_AGE,
                    'station': name,
                })
            nearby.sort()
            print(
                f"Air4Thai สถานีใน {WIDE_KM} กม.: "
                f"{[(f'{d:.1f}km', f'{a//60}m', n, v) for d,a,n,v in nearby[:5]]}"
            )
    except Exception as e:
        print(f"⚠️ Air4Thai error: {e}")

    WAQI_TOKEN = os.environ.get("WAQI_TOKEN")
    if WAQI_TOKEN:
        for sid in ["419585", "419584"]:
            try:
                res = requests.get(
                    f"https://api.waqi.info/feed/@{sid}/?token={WAQI_TOKEN}",
                    timeout=15)
                print(f"WAQI @{sid} HTTP: {res.status_code}")
                if res.status_code == 200:
                    d = res.json()
                    if d.get('status') == 'ok':
                        iaqi = d['data'].get('iaqi', {})
                        pm = _safe_float((iaqi.get('pm25') or {}).get('v'))
                        sname = d['data'].get('city', {}).get('name', sid)
                        print(f"WAQI @{sid} ({sname}) PM2.5: {pm}")
                        if pm is not None:
                            waqi_value = pm
                            break
            except Exception as e:
                print(f"⚠️ WAQI @{sid} error: {e}")

        if waqi_value is None:
            try:
                res = requests.get(
                    f"https://api.waqi.info/feed/geo:{PM_LAT};{PM_LON}/?token={WAQI_TOKEN}",
                    timeout=15)
                print(f"WAQI geo HTTP: {res.status_code}")
                if res.status_code == 200:
                    d = res.json()
                    if d.get('status') == 'ok':
                        geo = d['data'].get('city', {}).get('geo', [])
                        sname = d['data'].get('city', {}).get('name', '?')
                        if len(geo) == 2:
                            dist = get_dist(PM_LAT, PM_LON, float(geo[0]), float(geo[1]))
                            print(f"WAQI geo: {sname} ห่าง {dist:.1f} กม.")
                            if dist <= 60:
                                iaqi = d['data'].get('iaqi', {})
                                pm = _safe_float((iaqi.get('pm25') or {}).get('v'))
                                print(f"WAQI geo PM2.5: {pm}")
                                if pm is not None:
                                    waqi_value = pm
                            else:
                                print(f"⚠️ WAQI geo: ไกลเกิน ({dist:.1f} กม.) ข้าม")
            except Exception as e:
                print(f"⚠️ WAQI geo error: {e}")
    else:
        print("⚠️ WAQI_TOKEN ไม่พบ")

    OWM_API_KEY = os.environ.get("OWM_API_KEY")
    if OWM_API_KEY:
        try:
            res = requests.get(
                f"http://api.openweathermap.org/data/2.5/air_pollution"
                f"?lat={PM_LAT}&lon={PM_LON}&appid={OWM_API_KEY}",
                timeout=15)
            print(f"OWM HTTP: {res.status_code}")
            if res.status_code == 200:
                comp = res.json()['list'][0]['components']
                pm = _safe_float(comp.get('pm2_5'))
                print(f"OWM PM2.5: {pm}")
                if pm is not None:
                    owm_value = pm
        except Exception as e:
            print(f"⚠️ OWM error: {e}")
    else:
        print("⚠️ OWM_API_KEY ไม่พบ")

    try:
        res = requests.get(
            f"https://air-quality-api.open-meteo.com/v1/air-quality"
            f"?latitude={PM_LAT}&longitude={PM_LON}"
            f"&current=pm2_5&timezone=Asia%2FBangkok",
            headers=headers, timeout=20)
        res.raise_for_status()
        data = res.json()
        if 'current' in data:
            pm = _safe_float(data['current'].get('pm2_5'))
            print(f"Open-Meteo PM2.5: {pm}")
            if pm is not None:
                openmeteo_value = pm
    except Exception as e:
        print(f"⚠️ Open-Meteo error: {e}")


    strict = [r for r in air4thai_rows if r['distance'] <= STRICT_KM and r['age'] <= STRICT_AGE]
    if strict:
        strict.sort(key=lambda x: (x['distance'], x['age']))
        pm = _weighted_pm25(strict[:3])
        if pm is not None:
            selected_source = "Air4Thai ใกล้สุด ≤20 กม./≤1 ชม."
            print(f"✅ ใช้ {selected_source}: {pm:.1f}")
            return ({'pm25': f"{pm:.1f}", 'source': selected_source} if return_meta else f"{pm:.1f}")

    wide = [r for r in air4thai_rows if r['distance'] <= WIDE_KM and r['age'] <= WIDE_AGE]
    if wide:
        wide.sort(key=lambda x: (x['distance'], x['age']))
        pm = _weighted_pm25(wide[:3])
        if pm is not None:
            selected_source = "Air4Thai รอบกว้าง ≤50 กม./≤3 ชม."
            print(f"✅ ใช้ {selected_source}: {pm:.1f}")
            return ({'pm25': f"{pm:.1f}", 'source': selected_source} if return_meta else f"{pm:.1f}")

    if waqi_value is not None:
        selected_source = "WAQI สถานีใกล้เคียง"
        print(f"✅ ใช้ {selected_source}: {waqi_value:.1f}")
        return ({'pm25': f"{waqi_value:.1f}", 'source': selected_source} if return_meta else f"{waqi_value:.1f}")

    ceiling_candidates = [v for v in [gistda_value, openmeteo_value] if v is not None]
    if ceiling_candidates:
        pm = max(ceiling_candidates)
        selected_source = "เพดานคัดกรอง GISTDA/Open-Meteo"
        print(f"✅ ใช้ {selected_source}: {pm:.1f} จาก {ceiling_candidates}")
        return ({'pm25': f"{pm:.1f}", 'source': selected_source} if return_meta else f"{pm:.1f}")

    if owm_value is not None:
        selected_source = "OpenWeatherMap fallback"
        print(f"✅ ใช้ {selected_source}: {owm_value:.1f}")
        return ({'pm25': f"{owm_value:.1f}", 'source': selected_source} if return_meta else f"{owm_value:.1f}")

    stale = [r for r in air4thai_rows if r['distance'] <= WIDE_KM]
    if stale:
        stale.sort(key=lambda x: (x['age'], x['distance']))
        pm = _weighted_pm25(stale[:3])
        if pm is not None:
            selected_source = "Air4Thai stale ≤50 กม."
            print(f"⚠️ ใช้ {selected_source}: {pm:.1f}")
            return ({'pm25': f"{pm:.1f}", 'source': selected_source} if return_meta else f"{pm:.1f}")

    print("❌ ทุกแหล่งล้มเหลว → N/A")
    return ({'pm25': 'N/A', 'source': 'ทุกแหล่งล้มเหลว'} if return_meta else "N/A")

def get_weather():
    TOMORROW_API_KEY = os.environ.get("TOMORROW_API_KEY")
    temp, pm25, rain_prob, humidity, wind, uv = "N/A", "N/A", "N/A", "N/A", "N/A", "N/A"
    if TOMORROW_API_KEY:
        try:
            tmr_res = _tomorrow_forecast()
            if tmr_res:
                tl = tmr_res['timelines']
                current_data = (tl.get('minutely') or tl['hourly'])[0]['values']
                humidity     = round(current_data['humidity'], 1)
                wind         = round(current_data['windSpeed'], 1)
                rain_prob    = max(h['values']['precipitationProbability']
                                   for h in tmr_res['timelines']['hourly'][:12])
        except Exception: pass
    try:
        res  = _http_get_json(
            f"https://api.open-meteo.com/v1/forecast"
            f"?latitude={INBURI_LAT}&longitude={INBURI_LON}"
            f"&current=temperature_2m,uv_index&timezone=Asia%2FBangkok",
            timeout=15, retries=3)
        temp = res['current']['temperature_2m']
        uv   = res['current'].get('uv_index', 'N/A')
        cache_put("weather_om", {"temp": temp, "uv": uv})
    except Exception:
        cached, ts = cache_get("weather_om")
        if isinstance(cached, dict):
            temp, uv = cached.get("temp", "N/A"), cached.get("uv", "N/A")
            mark_stale("weather_om", ts)
            print(f"♻️ Open-Meteo ดึงสดไม่ได้ → ใช้ cache ณ {ts} น.")
    return temp, pm25, rain_prob, humidity, wind, uv

WATER_API_URL = os.environ.get("THAIWATER_WL_API_URL")   # (ไม่บังคับ) endpoint JSON ที่หา เจอจาก log "🔌 XHR"

def _fetch_water_api(names):
    """ยิง API JSON ตรง (เร็ว/ไม่พังเวลาหน้าเว็บเปลี่ยน DOM) → {ชื่อ: ระดับน้ำ หรือ None}
    ใช้เมื่อกำหนด THAIWATER_WL_API_URL เท่านั้น parser เป็นแบบ generic: หา dict ที่มีชื่อสถานี + ฟิลด์ระดับน้ำ"""
    found = {n: None for n in names}
    if not WATER_API_URL:
        return found
    data = _http_get_json(WATER_API_URL, timeout=20, retries=3)
    if not data:
        return found
    fields = ("waterlevel_msl", "water_level", "waterlevel", "wl", "level")

    def strings(o, depth=2):
        if isinstance(o, str):
            yield o
        elif isinstance(o, dict) and depth > 0:
            for v in o.values():
                yield from strings(v, depth - 1)

    hits = {n: [] for n in names}
    def walk(o):
        if isinstance(o, dict):
            val = None
            for f in fields:
                try:
                    v = float(str(o.get(f)).replace(",", ""))
                    if 0 < v < 100:
                        val = v
                        break
                except (TypeError, ValueError):
                    continue
            if val is not None:
                texts = list(strings(o))
                for n in names:
                    if any(n in t for t in texts):
                        hits[n].append((("สิงห์บุรี" in " ".join(texts)), val))
            for v in o.values():
                walk(v)
        elif isinstance(o, list):
            for v in o:
                walk(v)
    walk(data)
    for n in names:
        if hits[n]:
            hits[n].sort(key=lambda x: not x[0])      # เอาที่ระบุจังหวัดสิงห์บุรีก่อน
            found[n] = hits[n][0][1]
    print(f"🔌 Water API: {found}")
    return found

WL_MIN, WL_MAX = -5.0, 30.0     # ช่วงที่เป็นไปได้ของระดับน้ำ (ม.) ใช้ดักเลขผิดคอลัมน์

def _to_num(txt):
    c = re.sub(r"[^0-9.\-]", "", re.sub(r"[ ,]", "", txt or ""))
    if not c or c in ("-", ".", "-."):
        return None
    try:
        return float(c)
    except ValueError:
        return None

def _pick_level(heads, cells):
    """เลือกค่าระดับน้ำจากแถว: ใช้หัวคอลัมน์ที่มีคำว่า 'ระดับน้ำ' (ไม่ใช่ 'ตลิ่ง') ถ้าหาไม่เจอ
    ค่อยเดาคอลัมน์ตัวเลขแรก (พิมพ์เตือน) และทุกกรณีต้องอยู่ในช่วงที่เป็นไปได้ ไม่งั้นคืน None"""
    pos = None
    idx = next((i for i, h in enumerate(heads) if "ระดับน้ำ" in h and "ตลิ่ง" not in h), None)
    if idx is not None:
        pos = idx - (len(heads) - len(cells))        # หัวตารางรวมคอลัมน์ชื่อสถานี (th) ที่ไม่อยู่ใน cells
        if not (0 <= pos < len(cells)):
            pos = None
    if pos is not None:
        v = _to_num(cells[pos])
    else:
        print(f"⚠️ ไม่พบหัวคอลัมน์ 'ระดับน้ำ' ({heads}) → เดาคอลัมน์ตัวเลขแรก")
        v = next((n for n in (_to_num(c) for c in cells) if n is not None), None)
    if v is not None and not (WL_MIN <= v <= WL_MAX):
        print(f"⚠️ ระดับน้ำ {v} อยู่นอกช่วงที่เป็นไปได้ ({WL_MIN}..{WL_MAX}) → ทิ้ง")
        return None
    return v

def _scrape_water_once(names):
    """Playwright 1 รอบ (บล็อกรูป/ฟอนต์ให้โหลดเร็ว) + พิมพ์ URL ของ XHR JSON ลง log เพื่อหา API ตรง"""
    url = f"https://singburi.thaiwater.net/wl?cb={random.randint(10000, 99999)}"
    found = {n: None for n in names}
    seen, xhr = [], []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page    = browser.new_page()
        try:
            page.route("**/*", lambda route: route.abort()
                       if route.request.resource_type in ("image", "font", "media")
                       else route.continue_())

            def _on_resp(resp):
                try:
                    if ("json" in resp.headers.get("content-type", "")
                            and resp.request.resource_type in ("xhr", "fetch")):
                        xhr.append(resp.url.split("?")[0])
                except Exception:
                    pass
            page.on("response", _on_resp)

            page.goto(url, timeout=60000, wait_until="domcontentloaded")
            page.wait_for_selector("th[scope='row']", timeout=30000)
            soup = BeautifulSoup(page.content(), "html.parser")
            heads = [h.get_text(strip=True) for h in soup.select("thead th")]
            for th in soup.select("th[scope='row']"):
                label = th.get_text(strip=True)
                seen.append(label)
                for n in names:
                    if found[n] is not None or n not in label:
                        continue
                    cols = th.find_parent("tr").find_all("td")
                    found[n] = _pick_level(heads, [td.get_text(strip=True) for td in cols])
                    print(f"🔍 แถว '{label}' ดิบจากเว็บ: {[td.get_text(strip=True) for td in cols]} | หัวตาราง: {heads}")
        except Exception as e:
            print(f"เกิดข้อผิดพลาดในการดึงข้อมูลสิงห์บุรี: {e}")
        finally:
            browser.close()
    if xhr:
        print(f"🔌 XHR JSON ที่หน้าเว็บเรียก (ลองตั้ง THAIWATER_WL_API_URL ด้วยตัวที่คืนระดับน้ำ): {sorted(set(xhr))[:15]}")
    print(f"📋 สถานีที่เห็นบนหน้าเว็บ ({len(seen)}): {seen[:40]}")
    return found

def get_water_stations(names=("อินทร์บุรี", "โพนางดำ")):
    """ระดับน้ำหลายสถานี → {ชื่อสถานี: ระดับน้ำ หรือ None}
    ลำดับ: API ตรง (ถ้าตั้งไว้) → Playwright (2 ครั้ง) → cache ≤ 3 ชม. (ติดป้าย stale)"""
    found = _fetch_water_api(names)
    for attempt in range(2):
        missing = tuple(n for n in names if found[n] is None)
        if not missing:
            break
        if attempt:
            time.sleep(3 + random.uniform(0, 2))
        got = _scrape_water_once(missing)
        for n, v in got.items():
            if v is not None:
                found[n] = v

    for n in names:
        if found[n] is not None:
            cache_put(f"wl:{n}", found[n])
            continue
        cached, ts = cache_get(f"wl:{n}")
        if cached is not None:
            found[n] = cached
            mark_stale(f"wl:{n}", ts)
            print(f"♻️ ระดับน้ำ {n} ดึงสดไม่ได้ → ใช้ cache ณ {ts} น.")
    for n, v in found.items():
        tag = f" (cache ณ {_STALE[f'wl:{n}']} น.)" if f"wl:{n}" in _STALE else ""
        print(f"   {'✅' if v is not None else '❌ ไม่พบ'} {n}: {v}{tag}")
    return found

def get_inburi_data():
    """คงไว้เพื่อความเข้ากันได้กับโค้ดเดิม"""
    st = get_water_stations(("อินทร์บุรี",))
    return st.get("อินทร์บุรี"), ai_brain.BANK_LEVEL["อินทร์บุรี"]

def get_phonangdam_from_hii():
    """สำรอง: หาโพนางดำใน JSON ของ HII (หน้าเดียวกับที่ดึงเขื่อนเจ้าพระยา) รับเฉพาะฟิลด์ระดับน้ำที่ชัดเจน
    ถ้าไม่เจอจะพิมพ์โครงสร้างลง log เพื่อให้รู้ว่าต้องชี้ไปฟิลด์ไหน"""
    try:
        res = requests.get(
            f"https://tiwrm.hii.or.th/DATA/REPORT/php/chart/chaopraya/small/chaopraya.php"
            f"?cb={random.randint(10000, 99999)}", timeout=20)
        mt = re.search(r'var json_data = (\[.*\]);', res.text)
        if not mt:
            print("ℹ️ HII: ไม่พบ json_data")
            return None
        data = json.loads(mt.group(1))
        hits = []
        def walk(o, path):
            if isinstance(o, dict):
                if any(isinstance(v, str) and "โพนางดำ" in v for v in o.values()):
                    hits.append((path, o))
                for k, v in o.items():
                    walk(v, path + [k])
            elif isinstance(o, list):
                for i, v in enumerate(o):
                    walk(v, path + [i])
        walk(data, [])
        for path, o in hits:
            print(f"🔎 HII พบ 'โพนางดำ' ที่ {path}: { {k: o[k] for k in list(o)[:12]} }")
            for f in ("wl", "water_level", "level"):
                v = o.get(f)
                try:
                    return float(str(v).replace(",", ""))
                except Exception:
                    pass
        if not hits:
            iw = (data[0] or {}).get("itc_water") or {}
            print(f"ℹ️ HII ไม่พบชื่อ 'โพนางดำ' | รหัสสถานีใน itc_water: {list(iw.keys())[:60]}")
    except Exception as e:
        print(f"⚠️ HII โพนางดำ error: {e}")
    return None

def fetch_chao_phraya_dam_discharge():
    try:
        text = _http_get_text(
            f"https://tiwrm.hii.or.th/DATA/REPORT/php/chart/chaopraya/small/chaopraya.php"
            f"?cb={random.randint(10000, 99999)}", timeout=20, retries=3)
        match = re.search(r'var json_data = (\[.*\]);', text or "")
        if match:
            val = json.loads(match.group(1))[0]['itc_water']['C13']['storage']
            val = float(val) if isinstance(val, (int, float)) else float(str(val).replace(',', ''))
            cache_put("dam_discharge", val)
            return val
    except Exception as e:
        print(f"⚠️ เขื่อน error: {e}")
    cached, ts = cache_get("dam_discharge")
    if cached is not None:
        mark_stale("dam_discharge", ts)
        print(f"♻️ เขื่อนเจ้าพระยา ดึงสดไม่ได้ → ใช้ cache ณ {ts} น.")
        return float(cached)
    return None

# ─────────────────────────────────────────────
# ระบบบันทึกข้อมูลรายวัน (Save to CSV)
# ─────────────────────────────────────────────
def save_current_water_data(wl, discharge, pho_wl=None):
    csv_file = "history_water.csv"
    file_exists = os.path.isfile(csv_file)
    try:
        with open(csv_file, mode='a', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            if not file_exists:
                writer.writerow(["Date", "Time", "Station", "WaterLevel", "Discharge"])
            
            date_str_csv = now.strftime('%Y-%m-%d')
            time_str_csv = now.strftime('%H:%M')
            wl_val = wl if wl is not None else "-"
            dis_val = discharge if discharge is not None else "-"
            
            if wl is not None or discharge is not None:      # ไม่มีทั้งคู่ = ไม่เขียนแถว "-,-"
                writer.writerow([date_str_csv, time_str_csv, "อินทร์บุรี", wl_val, dis_val])
            if pho_wl is not None:
                writer.writerow([date_str_csv, time_str_csv, "โพนางดำ", pho_wl, "-"])
    except Exception as e:
        print(f"⚠️ บันทึกข้อมูลลงประวัติรายวันไม่ได้: {e}")

# ─────────────────────────────────────────────
# ดึงข้อมูลประวัติย้อนหลังของปีที่แล้ว (พร้อมคำนวณระยะห่างตลิ่ง)
# ─────────────────────────────────────────────
def get_historical_water_data(target_date):
    file_paths = ['ข้อมูลน้ำอินทร์บุรี2568.xlsx', 'โพนางดำ.xlsx']
    records = []
    
    try:
        target_last_year = target_date.replace(year=target_date.year - 1, tzinfo=None)
    except ValueError:
        target_last_year = target_date.replace(year=target_date.year - 1, day=target_date.day - 1, tzinfo=None)
        
    THAI_MONTHS_LOCAL = ["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
                   "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]
    
    station_data_pool = {"อินทร์บุรี": [], "โพนางดำ": []}
    
    # --- ส่วนที่ 1: ดึงจากไฟล์ Excel ---
    for file_path in file_paths:
        if not os.path.exists(file_path): continue
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
        except Exception: continue
            
        station_name = "อินทร์บุรี" if "อินทร์" in file_path else "โพนางดำ"
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            headers = [cell.value for cell in sheet[1]]
            
            date_col_idx, wl_col_idx, dis_col_idx = None, None, None
            for i, h in enumerate(headers):
                if not h: continue
                nh = str(h).replace(' ', '').replace('\n', '')
                if ('วันที่' in nh or 'วัน' in nh) and date_col_idx is None: date_col_idx = i
                elif nh.startswith('ระดับน้ำ') and wl_col_idx is None: wl_col_idx = i
                elif ('ปริมาณน้ำปล่อย' in nh or 'เขื่อน' in nh) and dis_col_idx is None: dis_col_idx = i
            
            if date_col_idx is None or wl_col_idx is None or dis_col_idx is None: continue
                
            for row in sheet.iter_rows(min_row=2, values_only=True):
                raw_date = row[date_col_idx]
                if not raw_date: continue
                
                dt = None
                if isinstance(raw_date, datetime): dt = raw_date
                elif isinstance(raw_date, str):
                    rd = str(raw_date).strip()
                    for fmt in ['%d/%m/%Y %H:%M', '%d/%m/%Y %H:%M:%S', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%d/%m/%Y']:
                        try:
                            dt = datetime.strptime(rd, fmt)
                            break
                        except Exception: pass
                    if not dt:
                        try: dt = datetime.strptime(rd.split()[0], '%Y-%m-%d')
                        except Exception:
                            try: dt = datetime.strptime(rd.split()[0], '%d/%m/%Y')
                            except Exception: pass
                            
                if dt and dt.year == target_last_year.year:
                    wl_val = str(row[wl_col_idx]).split('/')[0].strip()
                    dis_val = str(row[dis_col_idx]).replace('.0', '').strip()
                    if dis_val in ['None', 'nan', '']: dis_val = '-'
                    dt_naive = dt.replace(tzinfo=None)
                    
                    date_diff = abs((target_last_year.date() - dt_naive.date()).days)
                    dummy_target = datetime(2000, 1, 1, target_last_year.hour, target_last_year.minute)
                    dummy_record = datetime(2000, 1, 1, dt_naive.hour, dt_naive.minute)
                    time_diff_sec = abs((dummy_target - dummy_record).total_seconds())
                    
                    # --- คำนวณความห่างจากตลิ่ง ---
                    try:
                        wl_float = float(wl_val)
                        bank_lvl = 13.00 if station_name == "อินทร์บุรี" else 13.87
                        diff_bank = bank_lvl - wl_float
                        if diff_bank > 0:
                            b_status = f"ต่ำกว่าตลิ่ง {diff_bank:.2f} ม."
                        elif diff_bank < 0:
                            b_status = f"ล้นตลิ่ง {abs(diff_bank):.2f} ม."
                        else:
                            b_status = "พอดีระดับตลิ่ง"
                    except Exception:
                        b_status = "-"

                    station_data_pool[station_name].append({
                        'date_diff': date_diff, 'time_diff': time_diff_sec,
                        'dt': dt_naive, 'wl': wl_val, 'dis': dis_val, 'b_status': b_status
                    })

    # --- ส่วนที่ 2: ดึงจากไฟล์ CSV ---
    csv_file = "history_water.csv"
    if os.path.exists(csv_file):
        try:
            with open(csv_file, mode='r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                for row in reader:
                    d_str = row.get("Date", "").strip()
                    t_str = row.get("Time", "").strip()
                    s_name = row.get("Station", "").strip()
                    w_val = row.get("WaterLevel", "").strip()
                    d_val = row.get("Discharge", "").strip()
                    
                    if not d_str or s_name not in station_data_pool: continue
                    
                    try: dt = datetime.strptime(f"{d_str} {t_str}", "%Y-%m-%d %H:%M")
                    except Exception:
                        try: dt = datetime.strptime(d_str, "%Y-%m-%d")
                        except Exception: continue
                        
                    if dt.year == target_last_year.year:
                        date_diff = abs((target_last_year.date() - dt.date()).days)
                        dummy_target = datetime(2000, 1, 1, target_last_year.hour, target_last_year.minute)
                        dummy_record = datetime(2000, 1, 1, dt.hour, dt.minute)
                        time_diff_sec = abs((dummy_target - dummy_record).total_seconds())
                        
                        try:
                            wl_float = float(w_val)
                            bank_lvl = 13.00 if s_name == "อินทร์บุรี" else 13.87
                            diff_bank = bank_lvl - wl_float
                            if diff_bank > 0:
                                b_status = f"ต่ำกว่าตลิ่ง {diff_bank:.2f} ม."
                            elif diff_bank < 0:
                                b_status = f"ล้นตลิ่ง {abs(diff_bank):.2f} ม."
                            else:
                                b_status = "พอดีระดับตลิ่ง"
                        except Exception:
                            b_status = "-"
                            
                        station_data_pool[s_name].append({
                            'date_diff': date_diff, 'time_diff': time_diff_sec,
                            'dt': dt, 'wl': w_val, 'dis': d_val, 'b_status': b_status
                        })
        except Exception: pass

    # --- ส่วนที่ 3: เลือกข้อมูลที่ดีที่สุด ---
    for st_name, data_list in station_data_pool.items():
        if not data_list: continue
        
        min_date_diff = min(x['date_diff'] for x in data_list)
        closest_date_records = [x for x in data_list if x['date_diff'] == min_date_diff]
        best = min(closest_date_records, key=lambda x: x['time_diff'])
        b_dt = best['dt']
        
        thai_date_str = f"{b_dt.day} {THAI_MONTHS_LOCAL[b_dt.month - 1]} {b_dt.year + 543}"
        t_str = b_dt.strftime('%H:%M')
        
        if best['date_diff'] == 0:
             records.append(f"📌 {st_name} ปีที่แล้ว (ตรงกับวันนี้ วันที่ {thai_date_str} เวลา {t_str} น.) ระดับน้ำ {best['wl']} ม. ({best['b_status']}) | ระบายน้ำ {best['dis']} ลบ.ม./วินาที")
        else:
             records.append(f"📌 {st_name} ปีที่แล้ว (ข้อมูลใกล้เคียง วันที่ {thai_date_str} เวลา {t_str} น.) ระดับน้ำ {best['wl']} ม. ({best['b_status']}) | ระบายน้ำ {best['dis']} ลบ.ม./วินาที")

    if records:
        return "\n".join(records)
    return "ไม่มีบันทึกข้อมูลของปีที่แล้ว"


# ─────────────────────────────────────────────
# Main (เวอร์ชัน AI agent: ค้นข่าว -> วิเคราะห์ -> เขียน -> ตรวจ -> โพสต์)
# ─────────────────────────────────────────────
import ai_brain

FORCE_REVIEW = os.environ.get("FORCE_REVIEW", "0") == "1"          # 1 = ให้ AI ตรวจทุกรอบเหมือนเดิม
POST_MAX_CHARS = int(os.environ.get("POST_MAX_CHARS", "2500"))     # ปรับตามขีดจำกัดของแพลตฟอร์ม

def check_draft_rules(draft, facts, header, has_fire):
    """ตรวจร่างโพสต์ด้วยกฎง่าย ๆ คืนรายการปัญหา (ว่าง = ผ่าน ไม่ต้องเรียก AI ตรวจ)"""
    issues = []
    d = (draft or "").strip()
    if len(d) < 80:
        issues.append("ข้อความสั้นผิดปกติ")
    if len(d) > POST_MAX_CHARS:
        issues.append(f"ยาวเกิน {POST_MAX_CHARS} ตัวอักษร")
    if "สถานการณ์อินทร์บุรี" not in d:
        issues.append("ไม่มีหัวโพสต์")
    if re.search(r"(?<![A-Za-z0-9])(N/A|None|nan|null|undefined)(?![A-Za-z0-9])|\{[A-Za-z_]+\}", d):
        issues.append("มีค่าหลุดจากระบบ (N/A/None/placeholder)")
    if has_fire and not re.search(r"ไฟ|ความร้อน|hotspot", d, re.I):
        issues.append("พบจุดความร้อนแต่ไม่ได้กล่าวถึง")
    return issues

def minimal_post(facts, header):
    """โพสต์สำรองขั้นสุดท้าย (ไม่ใช้ AI) กรณี AI + template_post ล้มเหลวทั้งคู่"""
    w, pm, rs, wt = facts.get("weather", {}), facts.get("pm25", {}), facts.get("rain_storm", {}), facts.get("water", {})
    lines = [header,
             f"🌡️ อุณหภูมิ {w.get('temp_c', 'N/A')} °C | ความชื้น {w.get('humidity_pct', 'N/A')}%",
             f"🌫️ PM2.5 {pm.get('value', 'N/A')} µg/m³ ({pm.get('level', '-')})",
             f"🌧️ {rs.get('summary', '-')}"]
    if wt.get("risk_level"):
        lines.append(f"🌊 ความเสี่ยงน้ำ: {wt['risk_level']}")
    return "\n".join(lines)

DRY_RUN = os.environ.get("DRY_RUN", "0") == "1"                       # 1 = ไม่ส่ง webhook
POST_TEMPLATE_ON_AI_FAIL = os.environ.get("POST_TEMPLATE_ON_AI_FAIL", "1") != "0"

if __name__ == "__main__":
    print("=== เริ่มรวบรวมข้อมูลอินทร์บุรี ===")
    state = load_state()
    prev_wl, prev_discharge = state.get("last_water_level"), state.get("last_discharge")
    period = ai_brain.period_of_day(now.hour)

    # ── 1) เก็บข้อมูลดิบ ──
    temp, _, rain_prob, humidity, wind, uv = get_weather()
    pm25_meta = get_accurate_pm25(return_meta=True)
    pm25 = pm25_meta['pm25']
    stations = get_water_stations()
    wl, bank_level = stations.get("อินทร์บุรี"), ai_brain.BANK_LEVEL["อินทร์บุรี"]
    pho_wl = stations.get("โพนางดำ")
    if pho_wl is None or "wl:โพนางดำ" in _STALE:      # ค่าจาก cache → ลองแหล่งสำรอง HII สดก่อน
        _hii = get_phonangdam_from_hii()
        if _hii is not None:
            pho_wl = _hii
            cache_put("wl:โพนางดำ", _hii)
            _STALE.pop("wl:โพนางดำ", None)
    discharge = fetch_chao_phraya_dam_discharge()
    hotspots = get_hotspots()
    rain_info = get_comprehensive_rain_info()

    # ── 2) ประวัติ/บริบทน้ำ (โหลดก่อนบันทึกค่าใหม่) ──
    series = ai_brain.load_series()
    wl_live  = None if "wl:อินทร์บุรี" in _STALE else wl
    pho_live = None if "wl:โพนางดำ" in _STALE else pho_wl
    dis_live = None if "dam_discharge" in _STALE else discharge
    if any(v is not None for v in (wl_live, pho_live, dis_live)):
        save_current_water_data(wl_live, dis_live, pho_live)   # ไม่เขียนค่า cache ลงประวัติ เพราะไม่ใช่ค่า ณ เวลานี้
    water_ctx = ai_brain.build_water_context(now, wl, discharge, series, bank_level, pho_wl=pho_wl)
    if water_ctx.get("change_24h") is None and wl is not None and prev_wl is not None:
        water_ctx["change_since_last_run"] = round(float(wl) - float(prev_wl), 2)   # ไม่ใช่ "เมื่อวาน"
    risk = ai_brain.assess_water_risk(water_ctx)
    print(f"🌊 ความเสี่ยงน้ำ: {risk['level']} | {risk['reasons']}")

    new_state = dict(state)
    if wl_live is not None:  new_state["last_water_level"] = float(wl_live)
    if dis_live is not None: new_state["last_discharge"] = float(dis_live)
    save_state(new_state)

    # ── 3) รวมเป็น facts ก้อนเดียว ──
    tmr = rain_info.get('forecast') or {}
    tmd_obs = rain_info.get('tmd_obs') or {}
    rain_now_confirmed = bool(
        (tmr.get('rain_now_intensity') or 0) >= 0.5
        or (rain_info.get('actual_rain_1h') or 0) >= 1
        or (tmd_obs.get('rain_3h') or 0) >= 1)
    has_fire = isinstance(hotspots, int) and hotspots > 0
    facts = {
        "now": {"date": date_str, "time": time_str, "weekday": thai_day_of_week, "period": period},
        "weather": {"temp_c": temp, "uv": uv, "rain_prob_12h_pct": rain_prob,
                    "humidity_pct": humidity, "wind_ms": wind},
        "pm25": {"value": pm25, "source": pm25_meta.get('source'),
                 "level": classify_pm25_th(pm25)['label'],
                 "instruction": build_pm25_instruction(pm25, pm25_meta.get('source', ''))},
        "hotspots": {"count": hotspots if isinstance(hotspots, int) else None,
                     "status": "ระบบดาวเทียมขัดข้อง" if hotspots == "N/A" else "ok"},
        "rain_storm": {"summary": rain_info['summary'], "risk_level": rain_info['risk_level'],
                       "sources": rain_info.get('sources'), "rain_now_confirmed": rain_now_confirmed,
                       "max_rain_prob_24h_pct": tmr.get('max_rain_24h'),
                       "note": "ช่วงเวลาที่ตรงกับชั่วโมงปัจจุบันคือ 'ตอนนี้' ไม่ใช่อนาคต"},
        "water": {**water_ctx, "risk_level": risk['level'], "risk_reasons": risk['reasons']},
    }

    if _STALE:
        labels = {"tmd_obs": "ฝนสถานีอุตุฯ", "tmd_nwp": "พยากรณ์ NWP กรมอุตุฯ", "tomorrow": "พยากรณ์ Tomorrow.io",
                  "weather_om": "อุณหภูมิ/UV", "wl:อินทร์บุรี": "ระดับน้ำอินทร์บุรี",
                  "wl:โพนางดำ": "ระดับน้ำโพนางดำ", "dam_discharge": "ปริมาณน้ำปล่อยเขื่อนเจ้าพระยา"}
        facts["stale_data"] = {
            "note": "ข้อมูลต่อไปนี้ดึงสดไม่ได้ จึงใช้ค่าล่าสุดที่บันทึกไว้ ต้องระบุ 'ข้อมูล ณ เวลา' กำกับ "
                    "ห้ามเขียนเหมือนเป็นค่าปัจจุบัน",
            "items": {labels.get(k, k): f"ข้อมูล ณ {t} น." for k, t in _STALE.items()}}
        print(f"♻️ ใช้ข้อมูลจาก cache: {facts['stale_data']['items']}")

    facts = ai_brain.thai_dates(facts)   # วันที่ทั้งหมดเป็น วัน เดือน ปี(พ.ศ.) ก่อนส่งให้ AI/เทมเพลต

    header = f"**สถานการณ์อินทร์บุรี** (ข้อมูล ณ วัน{thai_day_of_week}ที่ {date_str} เวลา {time_str})"
    when_text = f"วัน{thai_day_of_week}ที่ {date_str} เวลา {time_str}"

    # ── 4) ค้นข่าวล่าสุด -> วิเคราะห์ -> เขียน -> ตรวจ ──
    prev_post = ai_brain.previous_post_brief(state)
    try:
        research_text, sources = ai_brain.research_latest(client, when_text, period, water_ctx, risk)
    except Exception as e:
        print(f"⚠️ ค้นข่าวไม่สำเร็จ ข้ามขั้นนี้: {e}")
        research_text, sources = "", []
    try:
        analysis = ai_brain.analyze(client, facts, risk, research_text, prev_post)
    except Exception as e:
        print(f"⚠️ analyze ไม่สำเร็จ ใช้ระดับความเสี่ยงจากกฎแทน: {e}")
        analysis = {"level": risk["level"], "focus": None}
    print(f"🧠 analysis: level={analysis.get('level')} focus={analysis.get('focus')}")

    final_post = ""
    try:
        draft = ai_brain.write_post(client, facts, analysis, research_text, prev_post, header, has_fire)
        issues = check_draft_rules(draft, facts, header, has_fire)
        if FORCE_REVIEW or issues:
            print(f"🔎 ส่งให้ AI ตรวจ/แก้ ({'บังคับ' if FORCE_REVIEW else issues})")
            final_post, leftover = ai_brain.review_and_fix(
                client, draft, facts, analysis, research_text, prev_post, header, has_fire)
            if leftover:
                print(f"ℹ️ ข้อสังเกตที่เหลือ: {leftover}")
        else:
            print("✅ ร่างผ่านกฎตรวจ ไม่ต้องเรียก AI ตรวจ (ประหยัด 1 call)")
            final_post = draft
    except Exception as e:
        print(f"❌ AI เขียนโพสต์ไม่สำเร็จ: {e}")
        if POST_TEMPLATE_ON_AI_FAIL:
            try:
                final_post = ai_brain.template_post(facts, analysis, header)
            except Exception as te:
                print(f"❌ template_post ล้มเหลว ใช้โพสต์สำรองขั้นต่ำ: {te}")
                final_post = minimal_post(facts, header)

    if final_post:
        final_post = final_post.strip() + "\n\n#อินทร์บุรีรอดมั้ย #VIIRS #GEE"
    print("\nข้อความที่จะโพสต์:\n", final_post)

    # ── 5) โพสต์ + จดจำ ──
    if DRY_RUN:
        print("\n🧪 DRY_RUN: ไม่ส่ง webhook / ไม่บันทึกความจำโพสต์")
    elif MAKE_WEBHOOK_URL and final_post:
        try:
            res = requests.post(MAKE_WEBHOOK_URL, json={"text_to_post": final_post}, timeout=30)
            if res.status_code == 200:
                print("\n✅ ส่ง Webhook สำเร็จ!")
                new_state = ai_brain.remember_post(new_state, final_post, analysis, facts, f"{date_str} {time_str}")
                save_state(new_state)
            else:
                print(f"\n❌ Webhook ล้มเหลว HTTP {res.status_code}")
        except Exception as e:
            print(f"\n❌ ส่ง Webhook ไม่สำเร็จ: {e}")
